# -*- coding: utf-8 -*-
"""仓库 Excel 取证脚本 D：完整长文本 + 出库vs字典名称匹配 + 私有原件结构核对（只读）"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from openpyxl import load_workbook

PATH = r"e:\匠星智造MES\docs\sample-data\材料物资表-出入库（系统字段）-脱敏版.xlsx"
wb = load_workbook(PATH, data_only=False)

print("### 完整长文本")
wi = wb["原材料入库"]
for addr in ["A1", "D2", "E2", "R2", "S2", "T2", "U2", "V2"]:
    v = wi[addr].value
    if v: print(f"  原材料入库!{addr} = {v!r}")
print("  周转物料入库!A1 =", repr(wb["周转物料入库"]["A1"].value))
print("  周转物料入库!T3 =", repr(wb["周转物料入库"]["T3"].value))
print("  委外加工入库 !A1 =", repr(wb["委外加工入库 "]["A1"].value))
print("  委外加工出库!A1 =", repr(wb["委外加工出库"]["A1"].value))
for cf in wi.conditional_formatting:
    for r in cf.rules:
        print("  COND_FMT 原材料入库", str(cf.sqref), r.type, r.operator if hasattr(r,'operator') else '', r.formula, "fill=", r.dxf.fill.bgColor.rgb if r.dxf and r.dxf.fill and r.dxf.fill.bgColor else None)

# 出库材料名称 是否在字典名称集合
wd = wb["原材料字典表"]
names = set()
for r in range(2, wd.max_row + 1):
    v = wd.cell(r, 2).value
    if v: names.add(str(v).strip())
wo = wb["原材料出库"]
print("\n### 出库名称 ∈ 字典名称?")
out_names = set()
for r in range(3, wo.max_row + 1):
    v = wo.cell(r, 4).value
    if v: out_names.add(str(v).strip())
for n in sorted(out_names):
    print(f"  {n!r}: {'在字典' if n in names else '不在字典(字典无此名称)'}")
# 入库名称
in_names = set()
for r in range(4, 17):
    v = wi.cell(r, 11).value
    if v: in_names.add(str(v).strip())
for n in sorted(in_names):
    print(f"  入库 {n!r}: {'在字典' if n in names else '不在字典'}")
print("  字典21个名称全list:", sorted(names))

# 私有原件结构核对（仅结构，不输出任何值）
print("\n### 私有原件结构核对（仅行列结构）")
try:
    wb2 = load_workbook(r"e:\匠星智造MES\资料\材料物资表-出入库（系统字段）.xlsx", data_only=True, read_only=True)
    for ws2 in wb2.worksheets:
        print(f"  {ws2.title!r}: {ws2.calculate_dimension()}")
    wb2.close()
except Exception as e:
    print("  私有原件读取失败(不阻塞):", e)
print("DONE_D")
