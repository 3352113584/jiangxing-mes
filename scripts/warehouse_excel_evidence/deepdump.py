# -*- coding: utf-8 -*-
"""仓库 Excel 取证脚本 B：行级转储 + 字典表分析 + 关联测试（只读，内部核对用）"""
import json, re, sys, io
from collections import Counter
from openpyxl import load_workbook

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

PATH = r"e:\匠星智造MES\docs\sample-data\材料物资表-出入库（系统字段）-脱敏版.xlsx"
wb = load_workbook(PATH, data_only=False)

def fmt(v, maxlen=28):
    if v is None: return ""
    if hasattr(v, "strftime"): return v.strftime("%Y-%m-%d")
    s = str(v).replace("\n", "⏎")
    return s[:maxlen]

print("### 工作簿级")
print("sheets:", [(ws.title, ws.sheet_state) for ws in wb.worksheets])
print("defined_names:", list(wb.defined_names.keys()) if hasattr(wb.defined_names, 'keys') else wb.defined_names)
try:
    print("creator:", wb.properties.creator, "| lastModifiedBy:", wb.properties.lastModifiedBy, "| modified:", wb.properties.modified)
except Exception as e:
    print("props err", e)

DUMP = {"原材料入库": (1, 21), "原材料出库": (1, 9), "周转物料入库": (1, 7),
        "周转物料出库": (1, 9), "委外加工入库 ": (1, 4), "委外加工出库": (1, 3)}

for ws in wb.worksheets:
    name = ws.title
    print("\n" + "=" * 80)
    print("### SHEET:", repr(name))
    if ws.merged_cells.ranges:
        print("merged:", [str(m) for m in ws.merged_cells.ranges][:25])
    for row in ws.iter_rows():
        for c in row:
            if c.comment:
                print(f"COMMENT {c.coordinate}: {str(c.comment.text)[:300]!r}")
    if ws.conditional_formatting:
        for cf in ws.conditional_formatting:
            print("COND_FMT:", str(cf.sqref), [r.type for r in cf.rules], [ (r.formula if hasattr(r,'formula') else None) for r in cf.rules])
    key = name.strip()
    if key in DUMP or name in DUMP:
        r1, r2 = DUMP.get(key, DUMP.get(name))
        for r in range(r1, min(r2, ws.max_row) + 1):
            cells = [fmt(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
            # 压缩尾部空列
            while cells and cells[-1] == "": cells.pop()
            print(f"R{r}: " + " ┃ ".join(cells))

# ---------- 字典表分析 ----------
print("\n" + "=" * 80)
print("### 原材料字典表分析")
ws = wb["原材料字典表"]
rows = []
for r in range(4, ws.max_row + 1):  # 表头在第3行？确认：A1商品类别 A2? 先dump前5行
    vals = [ws.cell(r, c).value for c in range(1, 5)]
    rows.append((r, vals))
# 先看前 6 行结构
for r, vals in rows[:6]:
    print("dictR", r, [fmt(v, 30) for v in vals])
hdr_row = None
for r, vals in rows:
    if vals[0] and "类别" in str(vals[0]):
        hdr_row = r
data = [(r, vals) for r, vals in rows if hdr_row and r > hdr_row and any(v not in (None, "") for v in vals)]
print("header_row=", hdr_row, "data_rows=", len(data))
cat = Counter(str(v[0]).strip() for r, v in data if v[0])
print("商品类别分布:", cat.most_common(10))
name_c = Counter(str(v[1]).strip() for r, v in data if v[1])
print("名称 top12:", name_c.most_common(12))
mat_c = Counter(str(v[3]).strip() for r, v in data if v[3])
print("材质种类数:", len(mat_c), "top10:", mat_c.most_common(10))
# 完全重复（类别+名称+规格+材质）
tup = Counter(tuple(str(x).strip() if x is not None else "" for x in v) for r, v in data)
dups = {k: c for k, c in tup.items() if c > 1}
print("四元组完全重复组数:", len(dups), "涉及行数:", sum(dups.values()))
for k, c in list(sorted(dups.items(), key=lambda x: -x[1]))[:8]:
    print("   DUP", c, [s[:25] for s in k])
# 名称+规格相同、材质不同（同义不同码维度）
ng = {}
for r, v in data:
    if v[1] and v[2]:
        ng.setdefault((str(v[1]).strip(), str(v[2]).strip()), set()).add(str(v[3]).strip() if v[3] else "")
same_nm_diff_mat = {k: m for k, m in ng.items() if len(m) > 1}
print("名称+规格相同但材质不同的组数:", len(same_nm_diff_mat))
for k, m in list(same_nm_diff_mat.items())[:5]:
    print("   NG", [s[:22] for s in k[0:1]], k[1][:22], "->", [s[:18] for s in m])
# 规格格式
pat = {"三段数字*d*d*": 0, "含字母前缀": 0, "中文描述": 0, "其他": 0}
for r, v in data:
    s = str(v[2]).strip() if v[2] else ""
    if not s: continue
    if re.fullmatch(r"\d+(\.\d+)?\*\d+(\.\d+)?\*\d+(\.\d+)?", s): pat["三段数字*d*d*"] += 1
    elif re.search(r"[\u4e00-\u9fa5]", s): pat["中文描述"] += 1
    elif re.search(r"[A-Za-z]", s): pat["含字母前缀"] += 1
    else: pat["其他"] += 1
print("规格格式分布:", pat)
mat_null = sum(1 for r, v in data if not v[3])
print("材质为空行数:", mat_null, "/", len(data))
spec_null = sum(1 for r, v in data if not v[2])
print("规格为空行数:", spec_null)

# ---------- 周转物料字典分析 ----------
print("\n### 周转物料字曲表分析")
ws = wb["周转物料字曲表"]
hdr = None
items = []
for r in range(1, ws.max_row + 1):
    a = ws.cell(r, 4).value
    if ws.cell(r, 5).value and hdr is None and "物料名称" in str(ws.cell(r, 5).value):
        hdr = r
        continue
    if hdr and r > hdr:
        vals = [ws.cell(r, c).value for c in range(2, 9)]
        if any(v not in (None, "") for v in vals):
            items.append((r, vals))
print("header_row=", hdr, "data_rows=", len(items))
code_c = Counter(str(v[2]).strip() for r, v in items if v[2] not in (None, ""))
code_dup = {k: c for k, c in code_c.items() if c > 1}
print("物料编号重复组数:", len(code_dup), "top:", sorted(code_dup.items(), key=lambda x: -x[1])[:6])
no_code = sum(1 for r, v in items if v[2] in (None, ""))
print("无物料编号行数:", no_code, "/", len(items))
ns = Counter((str(v[3]).strip(), str(v[4]).strip()) for r, v in items)
ns_dup = {k: c for k, c in ns.items() if c > 1}
print("名称+规格重复组数:", len(ns_dup), "涉及行数:", sum(ns_dup.values()))
for k, c in list(sorted(ns_dup.items(), key=lambda x: -x[1]))[:6]:
    print("   DUP", c, [s[:20] for s in k])
l2 = Counter(str(v[0]).strip() for r, v in items if v[0])
print("二级分类分布:", l2.most_common(10))

# ---------- 关联测试：出入库行 vs 原材料字典 ----------
print("\n### 关联测试（精确字符串匹配）")
wd = wb["原材料字典表"]
dict_set = set()
for r in range(hdr_row + 1 if hdr_row else 4, wd.max_row + 1):
    k = tuple(str(wd.cell(r, c).value).strip() if wd.cell(r, c).value is not None else "" for c in range(1, 5))
    if any(k):
        dict_set.add(k)
# 原材料入库：名称K、规格L、材质M、类别I
wi = wb["原材料入库"]
hit = miss = 0
for r in range(4, wi.max_row + 1):
    nm, sp, mt = wi.cell(r, 11).value, wi.cell(r, 12).value, wi.cell(r, 13).value
    if not nm: continue
    cat_i = str(wi.cell(r, 9).value or "").strip()
    k = (cat_i, str(nm).strip(), str(sp).strip(), str(mt).strip() if mt else "")
    # 字典无类别匹配的宽松版
    k2 = tuple([k[1], k[2], k[3]])
    d2 = {(t[1], t[2], t[3]) for t in dict_set}
    if k in dict_set or k2 in d2:
        hit += 1
    else:
        miss += 1
        print(f"  入库R{r} 未命中字典: {k[1][:14]}|{k[2][:20]}|{k[3][:12]}")
print(f"原材料入库 命中字典: {hit} 未命中: {miss}")
wo = wb["原材料出库"]
hit = miss = 0
d3 = {(t[1], t[2], t[3]) for t in dict_set}
for r in range(3, wo.max_row + 1):
    nm, sp, mt = wo.cell(r, 4).value, wo.cell(r, 5).value, wo.cell(r, 6).value
    if not nm: continue
    k2 = (str(nm).strip(), str(sp).strip(), str(mt).strip() if mt else "")
    if k2 in d3: hit += 1
    else:
        miss += 1
        print(f"  出库R{r} 未命中字典: {k2[0][:16]}|{k2[1][:20]}|{k2[2][:12]}")
print(f"原材料出库 命中字典: {hit} 未命中: {miss}")

# 出库重量一致性核算（理论重 vs 领用重量）
print("\n### 原材料出库 单件重量一致性")
for r in range(3, wo.max_row + 1):
    nm = wo.cell(r, 4).value
    if not nm: continue
    th, wd_, ln = wo.cell(r, 7).value, wo.cell(r, 8).value, wo.cell(r, 9).value
    q, w = wo.cell(r, 10).value, wo.cell(r, 11).value
    if isinstance(q, (int, float)) and isinstance(w, (int, float)) and q > 0:
        per = w / q
        theo = None
        if all(isinstance(x, (int, float)) for x in (th, wd_, ln)) and th and wd_ and ln:
            theo = th * wd_ * ln * 7.85 / 1e9  # kg? -> m*7.85t/m3 => T: mm^3*7.85e-12*1000? 直接算 T
            theo = th / 1000 * wd_ / 1000 * ln / 1000 * 7850 / 1000  # m*m*m*7850kg /1000 => T
        print(f"  R{r} {str(nm)[:10]} 规格{th}x{wd_}x{ln} 数量{q} 重量{w}T 单件{per:.4f}T 理论单件{('%.4f' % theo) if theo else 'NA'}T 偏差{('%+.2f%%' % ((per/theo-1)*100)) if theo else ''}")
print("DONE_B")
