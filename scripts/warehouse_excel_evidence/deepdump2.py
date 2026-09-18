# -*- coding: utf-8 -*-
"""仓库 Excel 取证脚本 C：字典表完整分析 + 定义名称 + 出库重量核算（只读）"""
import re, sys, io
from collections import Counter, defaultdict
from openpyxl import load_workbook

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
PATH = r"e:\匠星智造MES\docs\sample-data\材料物资表-出入库（系统字段）-脱敏版.xlsx"
wb = load_workbook(PATH, data_only=False)

print("### defined names")
for name, dn in wb.defined_names.items():
    print(f"  {name!r} -> {dn.value!r} dest={[str(d) for d in dn.destinations]}")

# ---------- 原材料字典表（表头在第1行，数据从第2行） ----------
ws = wb["原材料字典表"]
data = []
for r in range(2, ws.max_row + 1):
    v = [ws.cell(r, c).value for c in range(1, 5)]
    if any(x not in (None, "") for x in v):
        data.append((r, [str(x).strip() if x is not None else "" for x in v]))
print(f"\n### 原材料字典表 data_rows={len(data)}")
cat = Counter(v[0] for r, v in data)
print("商品类别分布:", cat.most_common(12))
name_c = Counter(v[1] for r, v in data)
print("名称去重数:", len(name_c), "top10:", name_c.most_common(10))
mat_c = Counter(v[3] for r, v in data if v[3])
print("材质去重数:", len(mat_c), "top8:", mat_c.most_common(8))
print("材质为空行数:", sum(1 for r, v in data if not v[3]))
print("规格为空行数:", sum(1 for r, v in data if not v[2]))
print("名称为空行数:", sum(1 for r, v in data if not v[1]))
tup = Counter(tuple(v) for r, v in data)
dups = {k: c for k, c in tup.items() if c > 1}
print("四元组(类别+名称+规格+材质)完全重复组数:", len(dups), "涉及行数:", sum(dups.values()))
for k, c in list(sorted(dups.items(), key=lambda x: -x[1]))[:6]:
    print("   DUP x%d: %s" % (c, " | ".join(s[:24] for s in k)))
ng = defaultdict(set)
for r, v in data:
    if v[1] and v[2]:
        ng[(v[1], v[2])].add(v[3])
multi = {k: m for k, m in ng.items() if len(m) > 1}
print("名称+规格相同但材质不同的组数:", len(multi))
for k, m in list(sorted(multi.items(), key=lambda x: -len(x[1])))[:5]:
    print("   NMGRP", k[0][:12], k[1][:22], "->", sorted(s[:16] for s in m))
pat = {"三段数字": 0, "四段数字": 0, "五段及以上": 0, "含中文": 0, "其他": 0}
for r, v in data:
    s = v[2]
    if not s: continue
    segs = s.split("*")
    if re.fullmatch(r"\d+(\.\d+)?(\*\d+(\.\d+)?){2}", s): pat["三段数字"] += 1
    elif re.fullmatch(r"\d+(\.\d+)?(\*\d+(\.\d+)?){3}", s): pat["四段数字"] += 1
    elif re.fullmatch(r"\d+(\.\d+)?(\*[0-9.]+){4,}.*", s): pat["五段及以上"] += 1
    elif re.search(r"[\u4e00-\u9fa5]", s): pat["含中文"] += 1
    else: pat["其他"] += 1
print("规格格式分布:", pat)
# 长度后缀"米"的规格
mi = [v[2] for r, v in data if v[2].endswith("米")]
print("以'米'结尾的规格行数:", len(mi), "示例:", mi[:4])

# ---------- 出库重量核算 ----------
print("\n### 原材料出库 领用重量 vs 理论重(7.85密度)")
wo = wb["原材料出库"]
for r in range(3, wo.max_row + 1):
    nm = wo.cell(r, 4).value
    if not nm: continue
    th, wdt, ln = wo.cell(r, 7).value, wo.cell(r, 8).value, wo.cell(r, 9).value
    q, w = wo.cell(r, 10).value, wo.cell(r, 11).value
    if isinstance(q, (int, float)) and isinstance(w, (int, float)) and q:
        per = w / q
        theo = th / 1000 * wdt / 1000 * ln / 1000 * 7850 / 1000 if all(isinstance(x, (int, float)) for x in (th, wdt, ln)) else None
        dev = f"{(per/theo-1)*100:+.2f}%" if theo else ""
        print(f"  R{r} {str(nm)[:12]} {th}x{wdt}x{ln} 件{q} 重{w}T 单件{per:.4f} 理论单件{theo and round(theo,4)} 偏差{dev}")

# ---------- 入库磅差核算 ----------
print("\n### 原材料入库 车次分组与磅差核算")
wi = wb["原材料入库"]
S = {r: wi.cell(r, 19).value for r in range(4, 17)}
grp1 = [S[r] for r in range(4, 11)]
ab1 = sum(grp1); ac1 = wi.cell(4, 29).value
print(f"  车次1(R4-R10合并): 货单合计={ab1:.3f}T 进厂磅重={ac1}T 磅差={(ab1-ac1)/ab1*1000:.2f}‰")
for r in range(11, 17):
    ab, ac = S[r], wi.cell(r, 29).value
    print(f"  R{r}: 领用单行重={ab}T 磅重={ac}T 磅差={(ab-ac)/ab*1000:+.2f}‰")

# ---------- 周转字典重复编号细节 ----------
print("\n### 周转字曲表 重复编号细节")
ws2 = wb["周转物料字曲表"]
rows2 = []
for r in range(3, ws2.max_row + 1):
    code = ws2.cell(r, 4).value
    nm = ws2.cell(r, 5).value
    spec = ws2.cell(r, 6).value
    unit = ws2.cell(r, 7).value
    if code not in (None, ""):
        rows2.append((str(code).strip(), str(nm or "").strip(), str(spec or "").strip(), str(unit or "").strip()))
code_rows = defaultdict(list)
for t in rows2:
    code_rows[t[0]].append(t)
dup_codes = {k: v for k, v in code_rows.items() if len(v) > 1}
print("重复编号组数:", len(dup_codes))
for k in ["YC0105", "FC3720", "FC3740"]:
    if k in dup_codes:
        names = Counter((t[1], t[2]) for t in dup_codes[k])
        print(f"  {k} x{len(dup_codes[k])}: 名称+规格分布={[(a[:14], b[:14], c) for (a, b), c in names.most_common(4)]}")
numeric_codes = [t for t in rows2 if t[0].isdigit()]
print("纯数字编号行数:", len(numeric_codes), "示例:", [t[0] for t in numeric_codes[:5]])

# ---------- 完整批注 ----------
print("\n### 全部批注（完整文本）")
for wsx in wb.worksheets:
    for row in wsx.iter_rows():
        for c in row:
            if c.comment:
                print(f"  [{wsx.title}!{c.coordinate}] 作者={c.comment.author!r} 文本={c.comment.text!r}")

# ---------- 周转字典编号列类型 ----------
ws3 = wb["周转物料字曲表"]
print("\n### 周转字曲表 表头区与前3行")
for r in range(1, 6):
    print(f"  R{r}:", [str(ws3.cell(r, c).value)[:16] if ws3.cell(r, c).value is not None else "" for c in range(1, 9)])
print("DONE_C")
