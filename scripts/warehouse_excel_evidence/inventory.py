# -*- coding: utf-8 -*-
"""仓库 Excel 取证脚本 A：8 Sheet 结构清点 + 全字段证据提取（只读）"""
import json, sys, io
from collections import Counter
from openpyxl import load_workbook

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

PATH = r"e:\匠星智造MES\docs\sample-data\材料物资表-出入库（系统字段）-脱敏版.xlsx"
OUT  = r"C:\Users\33521\AppData\Local\Temp\wh_evidence"

wb = load_workbook(PATH, data_only=False)   # 公式模式
wbv = load_workbook(PATH, data_only=True)   # 缓存值模式（脱敏版可能为 None）

def cell_str(v, maxlen=40):
    if v is None: return ""
    s = str(v).replace("\n", "\\n")
    return s[:maxlen]

report = {}
for ws in wb.worksheets:
    wsv = wbv[ws.title]
    info = {"dims": ws.dimensions, "max_row": ws.max_row, "max_col": ws.max_column,
            "freeze": str(ws.freeze_panes), "autofilter": str(ws.auto_filter.ref) if ws.auto_filter and ws.auto_filter.ref else None,
            "merged": [str(m) for m in ws.merged_cells.ranges],
            "hidden_cols": [], "hidden_rows_count": 0, "hidden_rows_sample": [],
            "data_validations": [], "comments": [], "defined_tables": list(ws.tables.keys()) if hasattr(ws, 'tables') else [],
            "cond_fmt_count": len(ws.conditional_formatting._cf_rules) if ws.conditional_formatting else 0,
            "columns": [], "empty_rows": [], "total_rows": [], "header_rows": []}
    # 隐藏列
    for letter, dim in ws.column_dimensions.items():
        if dim.hidden:
            info["hidden_cols"].append(letter)
    # 隐藏行
    hidden_rows = [r for r, d in ws.row_dimensions.items() if d.hidden]
    info["hidden_rows_count"] = len(hidden_rows)
    info["hidden_rows_sample"] = hidden_rows[:20]
    # 数据验证
    for dv in ws.data_validations.dataValidation:
        info["data_validations"].append({"ranges": str(dv.sqref), "type": dv.type,
                                         "formula1": cell_str(dv.formula1, 200) if dv.formula1 else None,
                                         "formula2": cell_str(dv.formula2, 100) if dv.formula2 else None,
                                         "allowBlank": dv.allow_blank, "prompt": cell_str(dv.prompt, 100) or None,
                                         "error": cell_str(dv.error, 100) or None})
    # 批注
    for row in ws.iter_rows():
        for c in row:
            if c.comment:
                info["comments"].append({"cell": c.coordinate, "text": cell_str(c.comment.text, 150)})
    # 逐列统计
    from openpyxl.utils import get_column_letter
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        vals, formulas = [], []
        formula_rows = []
        nonempty = 0
        num_ok, num_bad, dt_cnt = 0, 0, 0
        for r in range(1, ws.max_row + 1):
            v = ws.cell(r, col).value
            if v is None or (isinstance(v, str) and v.strip() == ""):
                continue
            nonempty += 1
            if isinstance(v, str) and v.startswith("="):
                formulas.append(v); formula_rows.append(r)
            elif isinstance(v, str):
                vals.append(v)
            else:
                vals.append(v)
                if hasattr(v, "year"): dt_cnt += 1
                elif isinstance(v, (int, float)): num_ok += 1
                else: num_bad += 1
        # 数字统计
        nums = [x for x in vals if isinstance(x, (int, float)) and not hasattr(x, "year")]
        num_stat = None
        if nums:
            num_stat = {"min": min(nums), "max": max(nums), "sum": round(sum(nums), 3)}
        # 样本值（最多 6 个去重，值转字符串截断）
        samples = []
        seen = set()
        for x in vals:
            k = cell_str(x, 40)
            if k and k not in seen:
                seen.add(k); samples.append(k)
            if len(samples) >= 6: break
        # 公式模式（去重，把行号换成 {r}）
        fpat = []
        fseen = set()
        for f in formulas[:2000]:
            import re
            p = re.sub(r"\d+", "{r}", f)
            if p not in fseen:
                fseen.add(p); fpat.append(p)
            if len(fpat) >= 8: break
        # 表头候选：前 5 行本列的值
        hdr = [cell_str(ws.cell(r, col).value, 30) for r in range(1, 6)]
        # 缓存值检查：公式列在值模式下是否有值
        cached = None
        if formula_rows:
            r0 = formula_rows[0]
            cached = cell_str(wsv.cell(r0, col).value, 30)
        info["columns"].append({"col": letter, "idx": col, "header_rows_1to5": hdr,
                                "nonempty": nonempty, "of_rows": ws.max_row,
                                "type_counts": {"text": sum(1 for x in vals if isinstance(x, str)),
                                                "num": num_ok, "datetime": dt_cnt},
                                "num_stat": num_stat, "samples": samples,
                                "formula_count": len(formulas), "formula_patterns": fpat,
                                "formula_first_row": formula_rows[0] if formula_rows else None,
                                "cached_value_available": cached})
    # 空行 / 合计行探测
    for r in range(1, ws.max_row + 1):
        row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        filled = [v for v in row_vals if v is not None and str(v).strip() != ""]
        if not filled:
            info["empty_rows"].append(r)
        else:
            joined = " ".join(str(v) for v in filled)
            if any(k in joined for k in ("合计", "总计", "小计", "SUM", "汇总")):
                info["total_rows"].append({"row": r, "content": joined[:120]})
    report[ws.title] = info

with open(OUT + r"\inventory.json", "w", encoding="utf-8") as f:
    json.dump(report, f, ensure_ascii=False, indent=1, default=str)

# 摘要输出
for name, info in report.items():
    print("=" * 70)
    print(f"SHEET: {name} | dims={info['dims']} rows={info['max_row']} cols={info['max_col']}")
    print(f"  freeze={info['freeze']} autofilter={info['autofilter']} merged={len(info['merged'])} hidden_cols={info['hidden_cols']} hidden_rows={info['hidden_rows_count']}")
    print(f"  dv_count={len(info['data_validations'])} comments={len(info['comments'])} cond_fmt={info['cond_fmt_count']} empty_rows={len(info['empty_rows'])} total_rows={len(info['total_rows'])}")
    if info["total_rows"]: print(f"  total_rows_detail={info['total_rows'][:3]}")
    if info["empty_rows"] and len(info["empty_rows"]) <= 15: print(f"  empty_rows_detail={info['empty_rows']}")
    for c in info["columns"]:
        hdr = c["header_rows_1to5"]
        hdr_s = " | ".join(h for h in hdr if h) or "(空)"
        print(f"  [{c['col']}] hdr5=[{hdr_s}] nonempty={c['nonempty']}/{c['of_rows']} types={c['type_counts']} formulas={c['formula_count']} samples={c['samples'][:4]}")
        if c["formula_patterns"]: print(f"        fml={c['formula_patterns'][:3]} first_row={c['formula_first_row']} cached={c['cached_value_available']}")
print("DONE")
